"""
migrar_ativos_planilha.py

Migra atribuições da planilha static/importacao_ativos.xlsx para o banco alvo.
IDs resolvidos dinamicamente — funciona em qualquer banco (local ou servidor).

Uso:
  python migrar_ativos_planilha.py          → dry-run (mostra o que fará, sem alterar)
  python migrar_ativos_planilha.py --apply  → aplica as mudanças
"""

import sys, os, sqlite3, unicodedata
from openpyxl import load_workbook
from datetime import datetime, timezone

def _norm(s):
    """Remove acentos e lowercaseeia."""
    return ''.join(
        c for c in unicodedata.normalize('NFD', s.lower())
        if unicodedata.category(c) != 'Mn'
    )

DRY_RUN = "--apply" not in sys.argv
os.environ.setdefault("PYTHONIOENCODING", "utf-8")
sys.stdout.reconfigure(encoding="utf-8")

BASE    = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE, "estoque.db")
XLSX    = os.path.join(BASE, "static", "importacao_ativos.xlsx")

# ── Alias de grupo: planilha_grupo → substrings aceitas no nome do grupo do banco
# Necessário pois servidor usa "Departamento X" e local pode usar "X"
GRUPO_ALIASES = {
    "comercial":     ["comercial", "sucesso do cliente"],
    "certitech":     ["certitech"],
    "fiscal":        ["fiscal", "administrativo"],   # Marcio e Lumara estão em ADM no servidor
    "atendimento":   ["atendimento", "administrativo"],
    "administrativo":["administrativo"],
    "processos":     ["processo"],
    "dp":            ["pessoal", "dp"],
    "rh":            ["recursos humanos", "rh"],
    "auditoria":     ["auditoria"],
    "contabil":      ["cont"],
    "financeiro":    ["financeiro", "administrativo"],  # Kamily está em ADM no servidor
}

# Correção de spelling (não de acento): planilha usa grafia diferente do servidor
PLANILHA_ATIVO_NOME_OVERRIDE = {
    "emilly":    "emily",
    "kamily":    "kamilly",
    "ana tereza": "ana teresa",
}

# Nome do novo ativo para casos sem ativo existente
NOVO_ATIVO_NOME = {
    ("eduardo",          "comercial"):    "Comercial - Eduardo",
    ("marcos",           "fiscal"):       "Fiscal - Marcos",
    ("elaine",           "atendimento"):  "ADM - Elaine",
    ("eduarda cipriano", "atendimento"):  "ADM - Eduarda Cipriano",
}

# Mapeamento de material da planilha para nome canônico do banco
# "Monitor" genérico permanece "Monitor" (sem tamanho)
PLANILHA_MAT_ALIAS = {
    "Mouse Sem Fio": "Mouse Sem Fio Recarregavel Wireless Led Rgb Ergonomico Longa Duracao",
    "Fone":          "Fone De Ouvido Headset Usb Telemarketing Com Microfone",
}

# Especificações para criar materiais que não existem
# grupo_nome + categoria_nome serão buscados/criados dinamicamente
NOVO_MAT_SPEC = {
    "Monitor":     {"grupo": "Monitor",    "cat": "Periferico",        "unidade": "un", "valor": 0.0},
    'Monitor 24"': {"grupo": "Monitor",    "cat": "Periferico",        "unidade": "un", "valor": 0.0},
    "Notebook":    {"grupo": "Notebook",   "cat": "Dispositivo Completo", "unidade": "un", "valor": 0.0},
    "Gabinete":    {"grupo": "Gabinete",   "cat": "Dispositivo",       "unidade": "un", "valor": 0.0},
    "All-in-One":  {"grupo": "All-in-One", "cat": "Dispositivo",       "unidade": "un", "valor": 0.0},
    "Teclado":     {"grupo": "Teclado",    "cat": "Periferico",        "unidade": "un", "valor": 0.0},
}

# ── helpers ───────────────────────────────────────────────────────────────────

def now_iso():
    return datetime.now(timezone.utc).isoformat()


def ler_planilha(xlsx_path):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Ativos"]
    itens = []
    for row in ws.iter_rows(min_row=2):
        v = [str(c.value or "").strip() for c in row[:13]]
        cat_at, grp_at, nome_at = v[0], v[1], v[2]
        nome_mat, cod = v[6], v[11]
        unidade = v[9] or "un"
        try:
            valor = float(v[10].replace(",", ".")) if v[10] else 0.0
        except ValueError:
            valor = 0.0
        if not nome_at and not nome_mat:
            break
        if nome_at and nome_mat:
            itens.append({
                "grp_ativo":  grp_at.strip(),
                "nome_ativo": nome_at.strip(),
                "nome_mat":   nome_mat.strip(),
                "unidade":    unidade,
                "valor":      valor,
                "codigo":     cod.strip(),
            })
    wb.close()
    return itens


# ── resolução de ativo ────────────────────────────────────────────────────────

def build_ativo_lookup(cur):
    cur.execute("""
        SELECT a.id, a.nome, ag.nome
        FROM ativos a JOIN ativos_grupos ag ON a.grupo_id = ag.id
        WHERE a.ativo = 1
    """)
    # Normaliza acentos para comparação robusta
    return [(r[0], _norm(r[1]), _norm(r[2])) for r in cur.fetchall()]


def grp_match(db_grp, planilha_grp):
    aliases = GRUPO_ALIASES.get(planilha_grp.lower(), [planilha_grp.lower()])
    return any(a in db_grp for a in aliases)


def resolve_ativo(cur, lookup, nome, grp, criados_cache, stats):
    """Retorna (ativo_id, criou). Procura no lookup; cria se necessário."""
    # Aplica override de spelling antes de normalizar acentos
    nome_corrigido = PLANILHA_ATIVO_NOME_OVERRIDE.get(nome.lower(), nome)
    nome_n = _norm(nome_corrigido)
    grp_n  = _norm(grp)
    chave  = (nome_n, grp_n)

    if chave in criados_cache:
        return criados_cache[chave], False

    # Estratégia 1: nome exato + grupo compatível
    matched = [r for r in lookup if r[1] == nome_n and grp_match(r[2], grp_n)]
    if len(matched) == 1:
        return matched[0][0], False

    # Estratégia 2: nome contido no nome do ativo (formato servidor "DEPT - Nome Sobrenome")
    matched = [r for r in lookup if nome_n in r[1] and grp_match(r[2], grp_n)]
    if len(matched) == 1:
        return matched[0][0], False

    # Estratégia 3: nome exato, sem restrição de grupo (Marcio, Kamily com depto diferente)
    matched = [r for r in lookup if r[1] == nome_n]
    if len(matched) == 1:
        return matched[0][0], False

    # Estratégia 4: nome contido, sem restrição de grupo
    nome_primeiro = nome_n.split()[0]
    if len(nome_primeiro) > 4:
        matched = [r for r in lookup if nome_primeiro in r[1]]
        if len(matched) == 1:
            return matched[0][0], False

    # Não encontrou — criar se tiver spec
    if chave not in NOVO_ATIVO_NOME:
        return None, False

    novo_nome = NOVO_ATIVO_NOME[chave]
    grupo_id  = _find_ativo_grupo(cur, grp)

    if grupo_id is None:
        print(f"  [ERRO] Grupo de ativo não encontrado para '{grp}'")
        return None, False

    if DRY_RUN:
        fake = -(len(criados_cache) + 1)
        criados_cache[chave] = fake
        stats.setdefault("ativos_criados", []).append(novo_nome)
        return fake, True

    cur.execute(
        "INSERT INTO ativos (nome, descricao, grupo_id, ativo, criado_em) VALUES (?,?,?,1,?)",
        (novo_nome, None, grupo_id, now_iso()),
    )
    new_id = cur.lastrowid
    criados_cache[chave] = new_id
    # Adicionar ao lookup para reutilização (já normalizado)
    lookup.append((new_id, _norm(novo_nome), _norm(grp)))
    stats.setdefault("ativos_criados", []).append(novo_nome)
    return new_id, True


def _find_ativo_grupo(cur, planilha_grp):
    aliases = GRUPO_ALIASES.get(planilha_grp.lower(), [planilha_grp.lower()])
    for alias in aliases:
        cur.execute(
            "SELECT id FROM ativos_grupos WHERE LOWER(nome) LIKE ?",
            (f"%{alias}%",),
        )
        r = cur.fetchone()
        if r:
            return r[0]
    return None


# ── resolução de material ─────────────────────────────────────────────────────

def build_mat_lookup(cur):
    cur.execute("SELECT LOWER(nome), id FROM materiais WHERE ativo=1")
    return dict(cur.fetchall())  # nome_lower → id


def resolve_material(cur, mat_lookup, nome_planilha, grupos_cache, stats):
    """Retorna (mat_id, criou)."""
    nome_real = PLANILHA_MAT_ALIAS.get(nome_planilha, nome_planilha)
    nome_l    = nome_real.lower()

    # Busca exata
    if nome_l in mat_lookup:
        return mat_lookup[nome_l], False

    # Busca por substring (ex: "Fone" dentro de "Fone De Ouvido Headset...")
    nome_lower_planilha = nome_planilha.lower()
    candidates = [mid for mname, mid in mat_lookup.items() if nome_lower_planilha in mname]
    if len(candidates) == 1:
        return candidates[0], False

    # Precisa criar
    spec = NOVO_MAT_SPEC.get(nome_real)
    if spec is None:
        return None, False

    grupo_id = _find_or_create_grupo_mat(cur, spec["grupo"], spec["cat"], grupos_cache)
    if grupo_id is None:
        print(f"  [ERRO] Grupo de material não encontrado para '{spec['grupo']}'")
        return None, False

    if DRY_RUN:
        fake = -(len(grupos_cache) + 200)
        mat_lookup[nome_l] = fake
        stats.setdefault("mats_criados", []).append(nome_real)
        return fake, True

    ts = now_iso()
    cur.execute(
        """INSERT INTO materiais
           (nome, descricao, quantidade, unidade, grupo_id, valor_unitario,
            fator_embalagem, usa_patrimonio, ativo, criado_em, atualizado_em)
           VALUES (?,?,0,?,?,?,1.0,1,1,?,?)""",
        (nome_real, None, spec["unidade"], grupo_id, spec["valor"], ts, ts),
    )
    new_id = cur.lastrowid
    mat_lookup[nome_l] = new_id
    stats.setdefault("mats_criados", []).append(nome_real)
    return new_id, True


def _find_or_create_grupo_mat(cur, nome_grupo, cat_substr, grupos_cache):
    if nome_grupo in grupos_cache:
        return grupos_cache[nome_grupo]

    # Buscar grupo existente
    cur.execute(
        "SELECT id FROM grupos_material WHERE LOWER(nome) LIKE ?",
        (f"%{nome_grupo.lower()}%",),
    )
    r = cur.fetchone()
    if r:
        grupos_cache[nome_grupo] = r[0]
        return r[0]

    # Buscar categoria
    cur.execute(
        "SELECT id FROM categorias WHERE LOWER(nome) LIKE ?",
        (f"%{cat_substr.lower()}%",),
    )
    r = cur.fetchone()
    if not r:
        return None
    cat_id = r[0]

    if DRY_RUN:
        fake = -(len(grupos_cache) + 500)
        grupos_cache[nome_grupo] = fake
        return fake

    cur.execute(
        "INSERT INTO grupos_material (nome, categoria_id, quantidade_minima, criado_em) VALUES (?,?,0.0,?)",
        (nome_grupo, cat_id, now_iso()),
    )
    new_id = cur.lastrowid
    grupos_cache[nome_grupo] = new_id
    return new_id


# ── atribuição ────────────────────────────────────────────────────────────────

def ja_atribuido(cur, ativo_id, mat_id):
    if ativo_id is None or ativo_id < 0:
        return False, None, None
    cur.execute(
        """SELECT ai.unidade_id, up.codigo
           FROM ativos_itens ai
           LEFT JOIN unidades_patrimonio up ON ai.unidade_id = up.id
           WHERE ai.ativo_id=? AND ai.material_id=? AND ai.devolvido_em IS NULL
           LIMIT 1""",
        (ativo_id, mat_id),
    )
    row = cur.fetchone()
    if row is None:
        return False, None, None
    return True, row[0], row[1]


def atualizar_codigo(cur, up_id, codigo):
    if DRY_RUN or not up_id or not codigo:
        return
    cur.execute(
        "UPDATE unidades_patrimonio SET codigo=? WHERE id=? AND (codigo IS NULL OR codigo='')",
        (codigo, up_id),
    )


def criar_atribuicao(cur, ativo_id, mat_id, codigo, stats):
    if DRY_RUN:
        stats["criados"] += 1
        return

    codigo_val = codigo if codigo else None

    cur.execute(
        """INSERT INTO unidades_patrimonio
           (material_id, status, origem, tag, codigo, criado_em)
           VALUES (?,?,?,?,?,?)""",
        (mat_id, "ativo", "importacao_planilha", "atribuido", codigo_val, now_iso()),
    )
    up_id = cur.lastrowid

    cur.execute(
        """INSERT INTO ativos_itens
           (ativo_id, material_id, unidade_id, quantidade, atribuido_em)
           VALUES (?,?,?,1.0,?)""",
        (ativo_id, mat_id, up_id, now_iso()),
    )

    # sync_qty: conta UPs ativas não-atribuídas
    cur.execute(
        """SELECT COUNT(*) FROM unidades_patrimonio
           WHERE material_id=? AND status='ativo'
           AND (tag IS NULL OR tag NOT IN ('atribuido','solicitado'))""",
        (mat_id,),
    )
    qty = cur.fetchone()[0]
    cur.execute("UPDATE materiais SET quantidade=? WHERE id=?", (qty, mat_id))

    stats["criados"] += 1


# ── main ──────────────────────────────────────────────────────────────────────

def main():
    print(f"{'[DRY-RUN] ' if DRY_RUN else ''}Banco: {DB_PATH}")
    print(f"Planilha: {XLSX}\n")

    itens = ler_planilha(XLSX)
    print(f"Itens na planilha: {len(itens)}\n")

    conn = sqlite3.connect(DB_PATH)
    cur  = conn.cursor()

    ativo_lookup  = build_ativo_lookup(cur)
    mat_lookup    = build_mat_lookup(cur)
    criados_atv   = {}
    grupos_cache  = {}

    stats = {"criados": 0, "pulados": 0, "sem_map": 0, "cod_atualizados": 0,
             "ativos_criados": [], "mats_criados": []}

    for item in itens:
        nome_at  = item["nome_ativo"]
        grp_at   = item["grp_ativo"]
        nome_mat = item["nome_mat"]
        codigo   = item["codigo"]

        # ── ativo ──
        ativo_id, ativo_novo = resolve_ativo(
            cur, ativo_lookup, nome_at, grp_at, criados_atv, stats
        )
        if ativo_id is None:
            print(f"  [SEM_MAP_ATIVO] {nome_at} / {grp_at} / {nome_mat}")
            stats["sem_map"] += 1
            continue

        # ── material ──
        mat_id, mat_novo = resolve_material(
            cur, mat_lookup, nome_mat, grupos_cache, stats
        )
        if mat_id is None:
            print(f"  [SEM_MAP_MAT] {nome_at} / {nome_mat}")
            stats["sem_map"] += 1
            continue

        # ── verificar duplicata ──
        existe, up_id, up_cod = ja_atribuido(cur, ativo_id, mat_id)
        if existe:
            if codigo and not up_cod:
                atualizar_codigo(cur, up_id, codigo)
                if DRY_RUN:
                    print(f"  ~ {nome_at:22s} | {nome_mat:30s} cod → {codigo}")
                stats["cod_atualizados"] += 1
            else:
                stats["pulados"] += 1
            continue

        # ── criar ──
        cod_info = f" cod={codigo}" if codigo else ""
        print(f"  + {nome_at:22s} | {nome_mat:45s}{cod_info}")
        criar_atribuicao(cur, ativo_id, mat_id, codigo, stats)

    if not DRY_RUN:
        conn.commit()
    conn.close()

    print("\n" + "=" * 60)
    print(f"Modo: {'DRY-RUN' if DRY_RUN else 'APLICADO'}")
    print(f"Criadas:          {stats['criados']}")
    print(f"Puladas (dupla):  {stats['pulados']}")
    print(f"Sem mapeamento:   {stats['sem_map']}")
    if stats["cod_atualizados"]:
        print(f"Cód. atualizados: {stats['cod_atualizados']}")
    if stats.get("ativos_criados"):
        print(f"Ativos criados:   {', '.join(stats['ativos_criados'])}")
    if stats.get("mats_criados"):
        print(f"Materiais criados:{', '.join(set(stats['mats_criados']))}")


if __name__ == "__main__":
    main()
