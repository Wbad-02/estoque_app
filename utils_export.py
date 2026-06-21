# © Wbad-02 — Todos os direitos reservados.
# github.com/Wbad-02
"""
Helpers genéricos para geração de relatórios Excel e PDF.
Centraliza estilização, headers, freeze panes e StreamingResponse.
"""
import io
from typing import Callable
from models import agora as _agora_br
from fastapi.responses import StreamingResponse

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm


HEADER_FILL = PatternFill("solid", fgColor="1E3A34")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ALERT_FILL = PatternFill("solid", fgColor="FFF3CD")
ZEBRA_COLOR = colors.HexColor("#F5F5F5")
GRID_COLOR = colors.HexColor("#CCCCCC")
BRAND_COLOR = colors.HexColor("#1E3A34")
COPYRIGHT = "© Todos os direitos reservados – github.com/Wbad-02"


def criar_excel(
    titulo_aba: str,
    headers: list[str],
    col_widths: list[int],
    rows: list[list],
    filename_prefix: str,
    row_style_fn: Callable | None = None,
    footer_row: list | None = None,
) -> StreamingResponse:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo_aba

    for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = w

    for ri, row in enumerate(rows, 2):
        styles = row_style_fn(ri - 2, row) if row_style_fn else {}
        fill = styles.get("fill")
        for ci, val in enumerate(row, 1):
            cell = ws.cell(row=ri, column=ci, value=val)
            align = styles.get(f"align_{ci}", styles.get("align", "left"))
            cell.alignment = Alignment(horizontal=align)
            fmt = styles.get(f"fmt_{ci}")
            if fmt:
                cell.number_format = fmt
            if fill:
                cell.fill = fill
            cell_font = styles.get(f"font_{ci}")
            if cell_font:
                cell.font = cell_font

    if footer_row:
        fri = len(rows) + 2
        for ci, val in enumerate(footer_row, 1):
            cell = ws.cell(row=fri, column=ci, value=val)
            if val:
                cell.font = Font(bold=True)

    ws.freeze_panes = "A2"
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"{filename_prefix}_{_agora_br().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def criar_pdf(
    titulo: str,
    headers: list[str],
    col_widths_cm: list[float],
    rows: list[list],
    filename_prefix: str,
    row_style_fn: Callable | None = None,
    footer_text: str | None = None,
    orientacao: str = "portrait",
) -> StreamingResponse:
    output = io.BytesIO()
    pagesize = landscape(A4) if orientacao == "landscape" else A4
    margins = {"leftMargin": 1*cm, "rightMargin": 1*cm} if orientacao == "landscape" else {"leftMargin": 1.5*cm, "rightMargin": 1.5*cm}
    doc = SimpleDocTemplate(output, pagesize=pagesize, **margins)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=14, textColor=BRAND_COLOR)
    sub_style = ParagraphStyle("s", parent=styles["Normal"], fontSize=9, textColor=colors.grey)
    cell_style = ParagraphStyle("c", parent=styles["Normal"], fontSize=8)

    elements = [
        Paragraph(titulo, title_style),
        Paragraph(f"Gerado em {_agora_br().strftime('%d/%m/%Y às %H:%M')}", sub_style),
        Spacer(1, 0.4 * cm),
    ]

    table_data = [headers]
    extra_styles = []
    for ri, row in enumerate(rows):
        pdf_row = []
        for val in row:
            if isinstance(val, str) and len(val) > 40:
                pdf_row.append(Paragraph(val, cell_style))
            else:
                pdf_row.append(val)
        table_data.append(pdf_row)

        if row_style_fn:
            rs = row_style_fn(ri, row)
            bg = rs.get("pdf_bg")
            if bg:
                extra_styles.append(("BACKGROUND", (0, ri + 1), (-1, ri + 1), colors.HexColor(bg)))
            txt_color = rs.get("pdf_text")
            if txt_color:
                extra_styles.append(("TEXTCOLOR", (0, ri + 1), (-1, ri + 1), colors.HexColor(txt_color)))

    cw = [w * cm for w in col_widths_cm]
    t = Table(table_data, colWidths=cw, repeatRows=1)
    base_styles = [
        ("BACKGROUND", (0, 0), (-1, 0), BRAND_COLOR),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTSIZE", (0, 1), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ZEBRA_COLOR]),
        ("GRID", (0, 0), (-1, -1), 0.3, GRID_COLOR),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]
    t.setStyle(TableStyle(base_styles + extra_styles))
    elements.append(t)
    elements.append(Spacer(1, 0.3 * cm))

    rodape = footer_text or f"Total: {len(rows)} registro(s) · {COPYRIGHT}"
    elements.append(Paragraph(rodape, sub_style))

    doc.build(elements)
    output.seek(0)
    filename = f"{filename_prefix}_{_agora_br().strftime('%Y%m%d_%H%M')}.pdf"
    return StreamingResponse(
        output, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
